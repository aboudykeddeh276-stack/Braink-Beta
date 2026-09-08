"""Theorem 1: Rule 110 algebraic normal form and spreadsheet-formula CA.

The ANF derivation `f(p, q, r) = q XOR r XOR (q AND r) XOR (p AND q AND r)`
is a correct, verifiable fact about the Rule 110 truth table (checked in
tests against the canonical table). What it does NOT establish is that a
real, deployed spreadsheet workbook is Turing-equivalent: Turing
completeness of Rule 110 (Cook, 2004) requires an unbounded tape and an
unbounded number of steps. A real `.xlsx` file has a finite grid and a
capped iterative-calculation depth, so any concrete workbook built from
this module is a bounded automaton, not a universal Turing machine. That
distinction is the whole point of exposing this as runnable code instead
of a claim in a document: `to_worksheet` below produces an actual,
finite, inspectable artifact, not a proof of universality.
"""

from __future__ import annotations

from openpyxl import Workbook
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

RULE_110_TABLE: dict[tuple[int, int, int], int] = {
    (1, 1, 1): 0,
    (1, 1, 0): 1,
    (1, 0, 1): 1,
    (1, 0, 0): 0,
    (0, 1, 1): 1,
    (0, 1, 0): 1,
    (0, 0, 1): 1,
    (0, 0, 0): 0,
}


def anf(p: int, q: int, r: int) -> int:
    """Rule 110's next-state function as an F2 polynomial: q XOR r XOR qr XOR pqr."""
    return q ^ r ^ (q & r) ^ (p & q & r)


def step(state: list[int]) -> list[int]:
    """Advance a 1D Rule 110 tape by one generation with fixed (0) boundaries."""
    padded = [0, *state, 0]
    return [anf(padded[i], padded[i + 1], padded[i + 2]) for i in range(len(state))]


def run(initial_state: list[int], generations: int) -> list[list[int]]:
    """Return `generations + 1` rows: the initial state followed by each step."""
    rows = [list(initial_state)]
    for _ in range(generations):
        rows.append(step(rows[-1]))
    return rows


def to_worksheet(ws: Worksheet, initial_state: list[int], generations: int) -> None:
    """Write the CA as live spreadsheet formulas (not pre-computed values).

    Row 1 holds the seed as literal 0/1 values. Each subsequent row is
    written as a MOD(...)-based formula referencing the row above, so
    opening the file and recalculating reproduces the same automaton —
    the formula is the mechanism, not a cached result. A one-column
    boundary of fixed zeros is kept on each side so every interior cell's
    formula can reference a real (left, center, right) triple without
    special-casing the tape edges.
    """
    width = len(initial_state) + 2  # + boundary columns on each side

    for col in range(1, width + 1):
        letter = get_column_letter(col)
        value = initial_state[col - 2] if 1 < col < width else 0
        ws[f"{letter}1"] = value

    for row in range(2, generations + 2):
        prev_row = row - 1
        for col in range(1, width + 1):
            letter = get_column_letter(col)
            if col in (1, width):
                ws[f"{letter}{row}"] = 0
                continue
            left = get_column_letter(col - 1)
            center = get_column_letter(col)
            right = get_column_letter(col + 1)
            formula = (
                f"=MOD({center}{prev_row}+{right}{prev_row}"
                f"+{center}{prev_row}*{right}{prev_row}"
                f"+{left}{prev_row}*{center}{prev_row}*{right}{prev_row},2)"
            )
            ws[f"{letter}{row}"] = formula


def to_workbook(initial_state: list[int], generations: int) -> Workbook:
    wb = Workbook()
    ws = wb.active
    ws.title = "rule110"
    to_worksheet(ws, initial_state, generations)
    return wb
